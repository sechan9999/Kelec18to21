"""k18to20_수정본.xlsx 에 data21 탭 추가 -> k18to21_수정본.xlsx
- 행: data20 과 같은 index 249개 (같은 순서, 같은 region·region2·district)
- 값: 21대 개표상황표 판독(ALL17.xlsx, 제외 행 빼고) 을 index 단위로 합산, 투표수·무효는 공개 21Data 투표구 값
- 비율 열은 data20 과 같은 수식, 20대 값·연령은 data20 에서 INDEX/MATCH 로 가져옴"""
import openpyxl, pandas as pd, numpy as np
from copy import copy
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

num = lambda s: pd.to_numeric(s.astype(str).str.replace(",", "").str.strip(), errors="coerce")

# Step 1. 21대 투표구 판독 + 공개 투표수·무효 (선거일·관내사전·관외사전·재외)
d = pd.read_excel("ALL17.xlsx", "투표구별"); d = d[d.판독방식 != "제외"].copy()
d["읍면동명"] = d.읍면동명.astype(str); d["투표구명"] = d.투표구명.astype(str).replace("nan", "")
P = pd.read_excel("Kelec18to21/corrected_data/K18to21charts_corrected.xlsx", "21Data", header=None).iloc[1:]
P = pd.DataFrame({"시도명": P[0], "구시군명": P[1], "읍면동명": P[2].astype(str), "투표구명": P[3].astype(str).replace("nan", ""),
                  "투표수": num(P[5]), "무효": num(P[11])}).dropna(subset=["투표수"])
x = d.merge(P, on=["시도명", "구시군명", "읍면동명", "투표구명"], how="left")
C5 = ["이재명", "김문수", "이준석", "권영국", "송진호"]
x["분류계"] = x[[f"분류_{c}" for c in C5]].sum(1)
print("투표수 결합 실패 행:", int(x.투표수.isna().sum()), "/", len(x))

# Step 2. 21대 이름 -> 18–20대 index 이름
code = pd.read_pickle("pe20.pkl").groupby("region").시도.first()
x["시도"] = x.시도명.replace({"세종특별자치시": "충청남도"})
MAP = {("경기도", "부천시원미구"): "부천시", ("경기도", "부천시소사구"): "부천시", ("경기도", "부천시오정구"): "부천시",
       ("경기도", "화성시갑"): "화성시", ("경기도", "화성시을"): "화성시", ("경기도", "여주시"): "여주군",
       ("충청북도", "청주시서원구"): "청주시흥덕구", ("충청북도", "청주시청원구"): "청원군",
       ("부산광역시", "부산진구"): "진구", ("인천광역시", "미추홀구"): "남구", ("충청남도", "세종특별자치시"): "세종시"}
x["district"] = [MAP.get((s, g), ("울산" + g) if s == "울산광역시" else g) for s, g in zip(x.시도, x.구시군명)]
x.loc[(x.시도명 == "대구광역시") & (x.구시군명 == "군위군"), "시도"] = "경상북도"
A = x.groupby(["시도", "district"]).agg(vote_all=("투표수", "sum"), 분류계=("분류계", "sum"), L1=("분류_이재명", "sum"), K1=("분류_김문수", "sum"),
                                       L2=("재확인_이재명", "sum"), K2=("재확인_김문수", "sum"), U2=("무효", "sum"),
                                       names21=("구시군명", lambda s: "·".join(sorted(set(s)))), 투표구=("구시군명", "size")).reset_index()
A["U_all"] = A.vote_all - A.분류계                       # 미분류(재확인) 전체 = 투표수 - 분류된 5후보 합 (무효 포함, data20 과 같은 정의)

# Step 3. data20 의 index 순서에 맞춤
wb = openpyxl.load_workbook("k18to20_수정본.xlsx")
w20 = wb["data20"]; h20 = {c.value: c.column for c in w20[1] if c.value}
rows20 = [dict(row=r, **{k: w20.cell(r, h20[k]).value for k in ["class", "region", "region2", "district", "index"]}) for r in range(2, 251)]
R20 = pd.DataFrame(rows20); R20["시도"] = R20.region.map(code)
M = R20.merge(A, on=["시도", "district"], how="left")
miss = M[M.L1.isna()]; print("21대 값 없는 index:", miss[["index", "시도", "district"]].values.tolist())
unused = A[~A.set_index(["시도", "district"]).index.isin(M.set_index(["시도", "district"]).index)]
print("index 에 안 들어간 21대 단위:", unused[["시도", "district"]].values.tolist())

# Step 4. data21 시트 쓰기 (data20 모양·수식 방식 그대로)
ws = wb.create_sheet("data21", index=wb.sheetnames.index("data20"))
HDR = ["class", "region", "region2", "district", "index", "vote_all", "U_all", "L1", "K1", "L2", "K2", "U2", "C_rate", "U_rate",
       "vote_all2", "U_all2", "U_rate2", "R1", "R2", "K", "over60", "R_1", "R_2", "K21_2", "R20_1", "R20_2", "K20_2", "district21", "n_precinct", "note"]
NOTE = {"L1": "이재명 분류표", "K1": "김문수 분류표", "L2": "이재명 재확인(미분류)", "K2": "김문수 재확인(미분류)", "U2": "무효 (공개 21Data)",
        "vote_all": "투표수 (공개 21Data, 선거일·관내사전·관외사전·재외 투표구 합; 거소·선상·잘못투입 제외)",
        "U_all": "투표수 − 분류된 5후보 합 (재확인 5후보 + 무효)", "over60": "20대(pe20res) 60세 이상 비율, data20 에서 가져옴",
        "R20_1": "data20 R_1 (같은 index)", "R20_2": "data20 R_2", "K20_2": "data20 K20_2", "district21": "합산한 21대 구시군 이름", "n_precinct": "합산한 투표구 수",
        "K21_2": "K = R_2/R_1 (김문수 분자)"}
hdr_src = {c.value: c for c in w20[1] if c.value}
for j, h in enumerate(HDR, 1):
    c = ws.cell(1, j, h); src = hdr_src.get(h) or hdr_src.get({"K1": "Y1", "K2": "Y2", "K21_2": "K20_2", "R20_1": "R19_1", "R20_2": "R19_2", "K20_2": "K19_2"}.get(h, ""), None)
    if src is not None: c.font = copy(src.font); c.fill = copy(src.fill); c.alignment = copy(src.alignment); c.border = copy(src.border)
    else: c.font = Font(name="Arial", bold=True, size=10)
    if h in NOTE: c.comment = Comment(NOTE[h], "data21", width=300, height=80)
body_font = copy(w20.cell(2, 1).font)
for i, r in M.iterrows():
    rr = i + 2; e = f"E{rr}"
    vals = [r["class"], r.region, r.region2, r.district, r["index"]]
    have = pd.notna(r.L1)
    vals += [int(r.vote_all), int(r.U_all), int(r.L1), int(r.K1), int(r.L2), int(r.K2), int(r.U2)] if have else [None] * 7
    f = (lambda s: s.replace("#", str(rr))) if have else (lambda s: None)
    vals += [f("=(F#-G#)/F#"), f("=G#/F#"), f("=F#-L#"), f("=G#-L#"), f("=P#/O#"), f("=I#/H#"), f("=K#/J#"), f("=S#/R#"),
             f"=INDEX(data20!$U$2:$U$250,MATCH({e},data20!$E$2:$E$250,0))",
             f("=I#/(I#+H#)"), f("=K#/(K#+J#)"), f("=W#/V#"),
             f"=INDEX(data20!$V$2:$V$250,MATCH({e},data20!$E$2:$E$250,0))", f"=INDEX(data20!$W$2:$W$250,MATCH({e},data20!$E$2:$E$250,0))",
             f"=INDEX(data20!$X$2:$X$250,MATCH({e},data20!$E$2:$E$250,0))",
             r.names21 if have else None, int(r.투표구) if have else None,
             None if have else "21대 판독 자료에 없음 (경북 개표상황표에 울릉군 없음)"]
    # 21대 행정구역 변화 메모
    notes = {"부천시": "2024년 원미·소사·오정 구 재도입 → 합산", "화성시": "화성시갑·을 합산", "청주시흥덕구": "흥덕구+서원구 (18–20대 행과 같은 구성)",
             "청원군": "청주시청원구", "군위군": "2023년 대구광역시로 편입, index 는 경북 유지", "남구": "미추홀구 (2018 개칭)" if r.region == 11 else None,
             "진구": "부산진구", "여주군": "여주시", "세종시": "세종특별자치시 (충남 코드 유지)"}
    if have and notes.get(r.district): vals[-1] = notes[r.district]
    for j, v in enumerate(vals, 1):
        c = ws.cell(rr, j, v); c.font = copy(body_font)
        if HDR[j - 1] in ("C_rate", "U_rate", "U_rate2", "over60", "R_1", "R_2", "R20_1", "R20_2"): c.number_format = "0.0000"
        elif HDR[j - 1] in ("R1", "R2", "K", "K21_2", "K20_2"): c.number_format = "0.0000"
        elif HDR[j - 1] in ("vote_all", "U_all", "L1", "K1", "L2", "K2", "U2", "vote_all2", "U_all2"): c.number_format = "#,##0"
    if not have:
        for j in range(1, len(HDR) + 1): ws.cell(rr, j).fill = PatternFill("solid", fgColor="FCD5B4")
for j, h in enumerate(HDR, 1): ws.column_dimensions[L(j)].width = 22 if h in ("district21", "note") else (12 if h == "district" else 10)
ws.column_dimensions["AD"].width = 44; ws.freeze_panes = "F2"

# Step 5. 점검 시트에 data21 설명 추가
au = wb["점검_수정내역"]; r0 = au.max_row + 2
lines = ["4. data21 탭 (21대, 2025) 추가",
         "행은 data20 과 같은 index 249개, 같은 순서·같은 district 이름. 21대 구시군을 index 에 맞춰 합산했습니다(부천 3개 구, 화성시갑·을, 청주 흥덕+서원 등). 합산한 21대 이름은 district21 열에 있습니다.",
         "값: 개표상황표 판독 결과(분류 = 분류된 투표지, 재확인 = 공개 최종득표 − 분류). 투표수·무효는 공개 21Data 투표구 값. 거소·선상·잘못투입 등 대조용 행은 제외.",
         "열 이름: L1·K1 = 이재명·김문수 분류, L2·K2 = 재확인(미분류), U2 = 무효. R_1·R_2·K21_2 는 김문수(보수 후보) 분자로 data20 의 R_1·R_2·K20_2 와 같은 정의.",
         "비율 열은 data20 과 같은 수식, over60·R20_1·R20_2·K20_2 는 data20 에서 index 로 가져오는 INDEX/MATCH 수식입니다.",
         f"경북 울릉군(index {int(miss['index'].iloc[0]) if len(miss) else '-'})은 21대 판독 자료에 없어 빈 행(주황색)입니다. 따라서 21대 n = 248 (이 탭 기준).",
         "요약 통계 행은 넣지 않았습니다 (SAS 가 관측치로 읽지 않도록)."]
for k, t in enumerate(lines):
    c = au.cell(r0 + k, 1, t); c.font = Font(name="Arial", bold=(k == 0), size=10 if k == 0 else 9); c.alignment = Alignment(wrap_text=False)
wb.save("k18to21_수정본.xlsx")
M.to_pickle("data21_units.pkl"); print("saved", len(M), "rows")
