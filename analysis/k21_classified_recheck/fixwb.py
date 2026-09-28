"""k18to20_08FEB23.xlsx 수정본 만들기
- 틀린 값만 고치고(노란색 + 메모), 공개값과 크게 다른 단위는 표시만(주황색 + 메모)
- 요약 통계 행을 data19/data20 에서 새 시트로 옮겨 SAS 가 관측치로 읽지 않게 함
- 새 시트 '점검_수정내역' 에 모든 변경·표시·재계산 결과 기록"""
import openpyxl, pandas as pd, numpy as np, statsmodels.formula.api as smf
from copy import copy
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

SRC, OUT = "k18to20.xlsx", "k18to20_수정본.xlsx"
wb = openpyxl.load_workbook(SRC)
YEL = PatternFill("solid", fgColor="FFFF00"); ORG = PatternFill("solid", fgColor="FCD5B4"); RED = PatternFill("solid", fgColor="F4B6B6")
F = "Arial"; AUTH = "점검"
log = []   # 구분, 시트, 셀, 구시군, 내용, 이전값, 이후값, 근거

def hdr(ws):
    return {(c.value.strip() if isinstance(c.value, str) else c.value): c.column for c in ws[1] if c.value is not None}
def row_of(ws, idx):
    col = hdr(ws)["index"]
    for r in range(2, 251):
        if ws.cell(r, col).value == idx: return r
    raise KeyError(idx)
def change(ws, cell, new, dist, what, why):
    c = ws[cell]; old = c.value; c.value = new; c.fill = YEL
    c.comment = Comment(f"[수정] {what}\n이전: {old}\n{why}", AUTH, width=320, height=110)
    log.append(dict(구분="수정", 시트=ws.title, 셀=cell, 구시군=dist, 내용=what, 이전값=old, 이후값=new, 근거=why))
def flag(ws, cell, dist, what, why, fill=ORG):
    c = ws[cell]; c.fill = fill
    c.comment = Comment(f"[확인 필요] {what}\n{why}", AUTH, width=320, height=110)
    log.append(dict(구분="표시", 시트=ws.title, 셀=cell, 구시군=dist, 내용=what, 이전값=c.value, 이후값=None, 근거=why))

d18, d19, d20, kc = wb["data18"], wb["data19"], wb["data20"], wb["Kcomparison"]
H18, H19, H20, HK = hdr(d18), hdr(d19), hdr(d20), hdr(kc)
cell = lambda H, name, r: f"{L(H[name])}{r}"

# 1) data19 index 31 이름 빈칸 -> 부천시 (다른 세 시트와 같게)
r = row_of(d19, 31)
change(d19, cell(H19, "district", r), "부천시", "부천시", "구시군 이름 빈칸 채움", "data18·data20·Kcomparison 의 같은 index(31) 이름이 부천시. 가나다순 위치·투표수 규모도 부천시와 일치")

# 2) data19 청주시흥덕구(index 250) 비율 열: 원자료로 다시 계산
r = row_of(d19, 250); v = {k: d19.cell(r, H19[k]).value for k in ["vote_all", "U_all", "M1", "H1", "D1", "M2", "H2", "D2", "U2"]}
exp = {"U_rate2": (v["U_all"] - v["U2"]) / (v["vote_all"] - v["U2"]), "C_rate": (v["vote_all"] - v["U_all"]) / v["vote_all"], "U_rate": v["U_all"] / v["vote_all"],
       "UM_rate": v["M2"] / (v["M1"] + v["M2"]), "UH_rate": v["H2"] / (v["H1"] + v["H2"]), "UD_rate": v["D2"] / (v["D1"] + v["D2"]), "K_DM": (v["D2"] / v["D1"]) / (v["M2"] / v["M1"])}
for k, e in exp.items():
    change(d19, cell(H19, k, r), e, "청주시흥덕구", f"{k} 재계산", "같은 행의 표 수(원자료와 일치)로 다시 계산한 값과 달라 교체. 표 수를 고친 뒤 비율을 갱신하지 않은 것으로 보임")

# 3) 요약 통계 행(254–258) 을 새 시트로 이동 — SAS 가 관측치로 읽어 19대 회귀 n=254 가 됨
st = wb.create_sheet("data19_요약통계", index=wb.sheetnames.index("data19") + 1)
st["A1"] = "data19 요약 통계 (원래 data19 254–258행)"; st["A1"].font = Font(name=F, bold=True)
st["A2"] = "SAS 가 이 행들을 관측치로 읽지 않도록 data19 에서 옮겼습니다. 수식은 data19!2:250 행을 그대로 참조합니다."; st["A2"].font = Font(name=F, italic=True, size=9)
for j in range(1, d19.max_column + 1):
    st.cell(4, j, d19.cell(1, j).value).font = Font(name=F, bold=True)
for k, r in enumerate(range(254, 259)):
    for j in range(1, d19.max_column + 1):
        src = d19.cell(r, j); val = src.value
        if isinstance(val, str) and val.startswith("="):
            val = val.replace("(", "(data19!", 1) if "(" in val else val        # =AVERAGE(G2:G250) -> =AVERAGE(data19!G2:G250)
        tgt = st.cell(5 + k, j, val); tgt.number_format = src.number_format; tgt.font = Font(name=F)
        src.value = None
log.append(dict(구분="수정", 시트="data19", 셀="A254:AF258", 구시군="(요약 통계 5행)", 내용="요약 통계 행을 data19_요약통계 시트로 이동",
                이전값="mean·sd·median·min·max 행 (R_1·R_2 포함)", 이후값="빈칸", 근거="SAS 19th 시트의 Observations Used 254 = 249 + 요약 5행. 이 행들이 회귀에 섞였음"))
kept = []
for r in range(254, 259):
    vals = [d20.cell(r, j).value for j in range(1, d20.max_column + 1)]
    if any(x is not None for x in vals):
        kept.append((r, vals))
        for j in range(1, d20.max_column + 1): d20.cell(r, j).value = None
log.append(dict(구분="수정", 시트="data20", 셀="Y254:AA258", 구시군="(요약 통계 5행)", 내용="19대 요약값 복사본 삭제 (data19_요약통계 와 같은 값)",
                이전값=f"{len(kept)}행", 이후값="빈칸", 근거="20th 시트의 19대 회귀(R19_2)에 섞일 수 있음. 같은 값은 data19_요약통계 시트에 있음"))

# 4) data20 오산시(index 51): 분류 열에 공개 최종득표가 들어 있음 -> 분류 = 최종 - 미분류
r = row_of(d20, 51); Y1, L1, Y2, L2 = (d20.cell(r, H20[k]).value for k in ["Y1", "L1", "Y2", "L2"])
change(d20, cell(H20, "Y1", r), Y1 - Y2, "오산시", "Y1(윤석열 분류) 복원", f"Y1 {Y1:,} = 공개 윤석열 최종득표와 정확히 같음. 분류 = 최종 − 미분류({Y2:,}). 단 Y2+L2+U2 가 U_all 보다 950표 많아 미분류 열도 확인 필요")
change(d20, cell(H20, "L1", r), L1 - L2, "오산시", "L1(이재명 분류) 복원", f"L1 {L1:,} = 공개 이재명 최종득표와 정확히 같음. 분류 = 최종 − 미분류({L2:,})")
# Kcomparison 의 오산 행 (data20 값을 숫자로 복사한 시트) 도 갱신
Y1n, L1n = Y1 - Y2, L1 - L2
R1, R2 = Y1n / L1n, Y2 / L2; R_1, R_2 = Y1n / (Y1n + L1n), Y2 / (Y2 + L2)
rk = row_of(kc, 51); kh = [c.value for c in kc[1]]
for j, name in enumerate(kh, 1):
    newv = {"R1": R1, "R2": R2, "R_1": R_1, "R_2": R_2}.get(name)
    if name == "K20": newv = R2 / R1
    if name == "K20_2": newv = R_2 / R_1
    if newv is not None:
        change(kc, f"{L(j)}{rk}", newv, "오산시", f"{name} 갱신", "data20 오산시 분류 열 복원에 맞춰 다시 계산")

# 5) 복원할 수 없는 오류·의심 값: 표시만
r = row_of(d20, 245)
for k in ["Y1", "L1"]:
    flag(d20, cell(H20, k, r), "제천시", f"{k} 오류 (복원 불가)", "Y1+L1 = 87,091 = 총투표수(vote_all). 분류·미분류 모두 공개값과 불일치. 분석에서 제외 권고", fill=RED)
r = row_of(d18, 31)
for k in ["vote_all", "under50", "r50s", "over60"]:
    flag(d18, cell(H18, k, r), "부천시", f"{k} 채워 넣은 값", "vote_all 500,000, 연령 0.64/0.20/0.15 는 딱 떨어지는 임의값으로 보임. 공개 투표수 517,155(원미·소사·오정 합). 표 수(P·M)는 정상 범위")

# 공개 최종득표와 차이 큰 단위 (18대 10% 이상, 19대 2% 이상, 20대 3% 이상·이름 문제 제외)
x18 = pd.read_pickle("pe18_out.pkl")["x"]; x19 = pd.read_pickle("pe19_out.pkl")["x"]
c20 = pd.read_pickle("pe20_check.pkl"); c20["rel"] = (c20.dY.abs() + c20.dL.abs()) / (c20.윤석열 + c20.이재명)
for ws, H, X, thr, cols, lab in [(d18, H18, x18, .10, ["P1", "M1"], "18대"), (d19, H19, x19, .02, ["M1", "H1"], "19대")]:
    for _, q in X[X.rel >= thr].iterrows():
        r = row_of(ws, int(q["index"]))
        flag(ws, cell(H, cols[0], r), q.district, f"공개 최종득표와 {q.rel:.0%} 차이", f"{lab} 공개 최종득표(보수+민주)보다 {q.rel:.1%} 적음. 투표구분 일부가 빠진 것으로 보임")
for _, q in c20[(c20.rel >= .03) & ~c20["index"].isin([51, 245, 250])].dropna(subset=["rel"]).iterrows():
    r = row_of(d20, int(q["index"]))
    flag(d20, cell(H20, "Y1", r), q.district, f"공개 최종득표와 {q.rel:.0%} 차이", f"20대 윤석열 {q.dY:+,.0f}표, 이재명 {q.dL:+,.0f}표 (분류+미분류 − 공개 최종)")
r = row_of(d20, 250)
flag(d20, cell(H20, "district", r), "청주시흥덕구", "서원구 포함", "표 수가 흥덕구+서원구 합과 비슷함(19대 자료와 같은 구성). 이름만 흥덕구")
flag(d19, cell(H19, "district", row_of(d19, 250)), "청주시흥덕구", "서원구 포함", "vote_all 287,210 ≈ 공개 흥덕구 153,064 + 서원구 134,980")

# 6) 점검 시트
au = wb.create_sheet("점검_수정내역", 0)
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def put(r, c, v, bold=False, fill=None, fmt=None, italic=False, size=10):
    x = au.cell(r, c, v); x.font = Font(name=F, bold=bold, italic=italic, size=size); x.alignment = Alignment(vertical="top", wrap_text=True)
    if fill: x.fill = fill
    if fmt: x.number_format = fmt
    return x
put(1, 1, "k18to20_08FEB23.xlsx 숫자 대조와 수정 내역", bold=True, size=13)
put(2, 1, "작성 2026-09-28. 노란색 = 고친 셀, 주황색 = 공개값과 차이 큰 셀(값은 그대로), 빨간색 = 복원할 수 없는 오류. 각 셀에 메모가 있습니다.", italic=True, size=9)
put(4, 1, "1. 대조 결과 요약", bold=True)
summ = [("원자료(pe18res2.csv, pe19res.xlsx, pe20res.xlsx) vs data18·19·20 표 수", "249행 × 7열 × 3선거 전부 일치"),
        ("data20 에 복사된 18·19대 값(R19_1…K18) vs data18·data19", "전부 일치 (index 기준)"),
        ("Kcomparison vs data20", "전부 일치 (오산시 행만 이번 수정으로 갱신)"),
        ("data18 R1·R2·K (숫자로 입력된 값) vs 원자료 재계산", "전부 일치"),
        ("data19 비율 열 vs 원자료 재계산", "청주시흥덕구 1행 7개 불일치 → 수정"),
        ("data19 254–258행 요약 통계", "SAS 가 관측치로 읽음 (19th 시트 n=254) → 새 시트로 이동"),
        ("공개 최종득표(중앙선관위) 대조", "18대 8곳 10% 이상, 19대 14곳 2% 이상, 20대 오산·제천 외 다수 차이 → 표시")]
put(5, 1, "항목", bold=True, fill=PatternFill("solid", fgColor="DDE4EE")); put(5, 2, "결과", bold=True, fill=PatternFill("solid", fgColor="DDE4EE"))
for i, (a, b) in enumerate(summ, 6): put(i, 1, a); put(i, 2, b)
r0 = 6 + len(summ) + 1
put(r0, 1, "2. SAS 결과와 재계산 비교 (R_2 = a + b·R_1, 보수 후보 분자)", bold=True)
p18 = pd.read_pickle("pe18_out.pkl")["p"]; p19 = pd.read_pickle("pe19_out.pkl")["p"]; p20 = pd.read_pickle("pe20.pkl").copy()
p20["R_1"] = p20.Y1 / (p20.Y1 + p20.L1); p20["R_2"] = p20.Y2 / (p20.Y2 + p20.L2)
q = p20.copy(); o = q["index"] == 51; q.loc[o, "Y1"] -= q.loc[o, "Y2"]; q.loc[o, "L1"] -= q.loc[o, "L2"]; q = q[q["index"] != 245]
q["R_1"] = q.Y1 / (q.Y1 + q.L1); q["R_2"] = q.Y2 / (q.Y2 + q.L2)
rows = [("18대 선형", "18th", 249, 0.9823, None, p18), ("19대 선형", "19th (요약 5행 포함)", 254, 0.9696, 0.04095, p19), ("19대 2차", "19th (요약 5행 포함)", 254, 0.9881, 0.02565, p19),
        ("20대 선형", "20th", 249, 0.9833, 0.02832, p20), ("20대 2차", "20th", 249, 0.9863, 0.02567, p20), ("20대 선형 (오산 복원·제천 제외)", "수정본 기준", None, None, None, q)]
hd = ["모형", "SAS 시트", "SAS n", "SAS R²", "SAS Root MSE", "재계산 n", "절편", "기울기", "재계산 R²", "재계산 Root MSE"]
for j, h in enumerate(hd, 1): put(r0 + 1, j, h, bold=True, fill=PatternFill("solid", fgColor="DDE4EE"))
for i, (nm, sh, n, r2, rm, d) in enumerate(rows, r0 + 2):
    d = d.assign(R_1sq=d.R_1 ** 2); m = smf.ols("R_2 ~ R_1sq + R_1" if "2차" in nm else "R_2 ~ R_1", d).fit()
    for j, v in enumerate([nm, sh, n, r2, rm, int(m.nobs), m.params.Intercept, m.params.R_1, m.rsquared, np.sqrt(m.mse_resid)], 1):
        put(i, j, v, fmt="0.0000" if isinstance(v, float) else None)
r1 = r0 + 2 + len(rows) + 1
put(r1, 1, "19th 시트 앞쪽 두 회귀는 요약 통계 5행이 섞인 결과입니다. 재계산(n 249) 값을 쓰십시오. 18th·20th 시트 회귀는 재계산과 같습니다.", italic=True, size=9)
put(r1 + 2, 1, "3. 변경·표시 목록", bold=True)
cols = ["구분", "시트", "셀", "구시군", "내용", "이전값", "이후값", "근거"]
for j, h in enumerate(cols, 1): put(r1 + 3, j, h, bold=True, fill=PatternFill("solid", fgColor="DDE4EE"))
for i, e in enumerate(log, r1 + 4):
    for j, h in enumerate(cols, 1):
        v = e[h]; v = v if v is None or isinstance(v, (int, float, str)) else str(v)
        put(i, j, v, fill=(YEL if e["구분"] == "수정" else ORG) if j == 1 else None, fmt="#,##0.######" if isinstance(v, float) else None)
for c, w in zip("ABCDEFGHIJ", [30, 22, 12, 14, 34, 18, 18, 70, 12, 14]): au.column_dimensions[c].width = w
au.freeze_panes = "A6"
wb.save(OUT)
print("수정", sum(e["구분"] == "수정" for e in log), "표시", sum(e["구분"] == "표시" for e in log), "->", OUT)
pd.DataFrame(log).to_pickle("fixwb_log.pkl")
