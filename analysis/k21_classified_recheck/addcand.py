"""k18to21_수정본.xlsx 에 후보별K 탭 추가 (18·19·21대 후보별 K, 21대 유형별, 60대 이상 비율과의 관계)"""
import shutil, openpyxl, pandas as pd, numpy as np
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
F = "k18to21_수정본.xlsx"; shutil.copy(F, "k18to21_수정본_before_cand2.xlsx")
T = pd.read_pickle("candk.pkl"); TY = pd.read_pickle("candk21_type.pkl"); A = pd.read_pickle("candage.pkl")
wb = openpyxl.load_workbook(F)
if "후보별K" in wb.sheetnames: del wb["후보별K"]
ws = wb.create_sheet("후보별K", wb.sheetnames.index("K정규성") + 1)
B = Font(bold=True); HF = PatternFill("solid", fgColor="DDEBF7"); RED = Font(bold=True, color="C00000"); BLU = Font(bold=True, color="1F4E99")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def hdr(r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = B; c.fill = HF; c.border = BOX; c.alignment = Alignment(horizontal="center", wrap_text=True)
def row(r, vals, fmt="0.000"):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.border = BOX
        if isinstance(v, float): c.number_format = fmt
r = 1
ws.cell(r, 1, "후보별 K — 18·19·21대 (20대는 후보별 자료 없음)").font = Font(bold=True, size=14); r += 1
ws.cell(r, 1, "K = (미분류 후보/민주 후보) ÷ (분류 후보/민주 후보). 18대 뉴스타파 251곳, 19대 선관위 250곳(봉화 개표단위 합계), 21대 개표상황표 판독 252곳"); r += 2

ws.cell(r, 1, "1. 후보별 K").font = B; r += 1
CAMP = {"박근혜": "보수", "홍준표": "보수", "김문수": "보수", "유승민": "보수", "이준석": "보수", "안철수": "중도", "심상정": "진보", "권영국": "진보", "기타": "군소", "송진호": "군소"}
hdr(r, ["선거", "후보", "성향", "기준", "분류표 득표율", "미분류표 득표율", "K 전국 합산", "K 구·시·군 평균", "K 중앙값", "K > 1 인 곳", "n"]); r += 1
for _, t in T.iterrows():
    row(r, [t.선거, t.후보, CAMP[t.후보], t.기준, float(t.분류득표율), float(t.미분류득표율), float(t.K_전국), float(t.K_구시군평균), float(t.K_중앙값), int(t.K_1초과), int(t.n)])
    for j in (5, 6): ws.cell(r, j).number_format = "0.00%"
    ws.cell(r, 7).font = RED if t.K_전국 > 1 else BLU; r += 1
r += 1
ws.cell(r, 1, "2. 21대 투표 유형별 K (이재명 대비, 전국 합산)").font = B; r += 1
hdr(r, ["유형", "김문수", "이준석", "권영국", "송진호"]); r += 1
for _, t in TY.iterrows(): row(r, [t.유형, float(t.김문수), float(t.이준석), float(t.권영국), float(t.송진호)]); r += 1
ws.cell(r, 1, "재외는 표 수가 적어 참고용."); r += 2

ws.cell(r, 1, "3. 후보별 log K 와 60대 이상 비율 (구·시·군, HC3)").font = B; r += 1
hdr(r, ["선거", "후보", "기준", "n", "상관", "60대 이상 10%p 당 log K 변화", "t", "R1 통제 후 변화", "t (R1 통제)"]); r += 1
for _, t in A.iterrows():
    row(r, [t.선거, {"박": "박근혜", "홍": "홍준표", "안": "안철수", "유": "유승민", "심": "심상정"}.get(t.후보, t.후보), {"문": "문재인"}.get(t.기준, t.기준), int(t.n),
            float(t.상관), float(t.기울기_10pt), float(t.t), float(t.기울기_R1통제_10pt), float(t.t_R1통제)]); r += 1
ws.cell(r, 1, "60대 이상 비율: 18대 뉴스타파, 19대 data19, 21대 data21(20대 기준 값). R1 = 분류표의 보수 주 후보 비율."); r += 2

ws.cell(r, 1, "4. 결론").font = B; r += 1
for t in ["고령층 지지가 두터운 보수 후보는 K > 1 (박근혜 1.39, 홍준표 1.61, 김문수 1.23).",
          "젊은 층 지지가 두터운 보수 후보는 K ≤ 1 (유승민 0.92, 이준석 0.97). 진보 후보도 K < 1 (심상정 0.74, 권영국 0.93).",
          "군소 후보는 세 선거 모두 K 가 가장 크다 (18대 기타 4.72, 19대 기타 2.56, 21대 송진호 3.34).",
          "군소 후보 K 는 60대 이상 비율이 높은 구·시·군일수록 커진다 (세 선거 모두 t 6–9).",
          "주 후보 K 는 고령 지역에서 커지지 않는다. 19대는 오히려 작아진다. 고령 지역은 두 후보 지지층의 연령 차이가 작아서일 수 있다(구·시·군 자료로는 구분 불가).",
          "치우침은 '보수 후보에게 몰아주기'보다 지지층의 표기 습관과 분류기 판독 특성으로 설명된다. 개인 단위 확인에는 투표구·투표지 자료가 필요하다."]:
    ws.cell(r, 1, "• " + t); r += 1
r += 1
EX = pd.read_pickle("explain.pkl")
ws.cell(r, 1, "5. K 를 설명하는 두 사실 (그림)").font = B; r += 1
hdr(r, ["선거", "후보", "득표율–60대 이상 상관", "K 전국 합산", "후보 미분류율"]); r += 1
for _, t in EX["R"].iterrows():
    row(r, [t.선거, t.후보, float(t.고령상관), float(t.K), float(t.미분류율)]); ws.cell(r, 5).number_format = "0.00%"; r += 1
ws.cell(r, 1, f"후보 11명: 득표율–60대 이상 상관과 log K 의 상관 {np.corrcoef(EX['R'].고령상관, EX['R'].logK)[0, 1]:.2f}"); r += 2
hdr(r, ["선거", "구·시·군 미분류율과 60대 이상 비율 상관", "60대 이상 10%p 당 미분류율 변화(%p)"]); r += 1
for e, d in EX["dist"].items():
    row(r, [e, float(np.corrcoef(d.age, d.urate)[0, 1]), float(np.polyfit(d.age, d.urate, 1)[0] * 10)]); r += 1
r += 1
from openpyxl.drawing.image import Image
img = Image("K_설명_고령지지.png"); img.width, img.height = 1200, 540; ws.add_image(img, f"A{r}")
ws.column_dimensions["A"].width = 12; ws.column_dimensions["B"].width = 12
for j in range(3, 12): ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 15
lg = wb["점검_수정내역"]; k = next((c.row for c in lg["A"] if isinstance(c.value, str) and c.value.startswith("후보별K:")), lg.max_row + 1)
lg.cell(k, 1, "후보별K: 18·19·21대 후보별 K(민주 후보 대비), 21대 유형별, 60대 이상 비율과의 관계, 설명 그림(후보 11명 고령 상관–K 상관 0.92). 유승민 0.92·이준석 0.97 로 보수 후보라도 K ≤ 1, 군소 후보 K 2.6–4.7 이 고령 지역에서 커짐.")
wb.save(F); print("saved", r)
