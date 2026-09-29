"""k18to21_수정본.xlsx 에 nec19 탭 추가: 19대 선관위 분류기 통계 원자료(구·시·군 250행) + data19 대조 + 후보별 재확인/분류 비교"""
import shutil, openpyxl, pandas as pd, numpy as np
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
F = "k18to21_수정본.xlsx"; shutil.copy(F, "k18to21_수정본_before_nec19.xlsx")
g = pd.read_pickle("nec19.pkl"); m = pd.read_pickle("nec19_merge.pkl"); x = pd.read_pickle("pe19_out.pkl")["x"]
wb = openpyxl.load_workbook(F)
if "nec19" in wb.sheetnames: del wb["nec19"]
ws = wb.create_sheet("nec19", wb.sheetnames.index("data19") + 1)
B = Font(bold=True); H = PatternFill("solid", fgColor="DDEBF7"); Y = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def hdr(r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = B; c.fill = H; c.border = BOX; c.alignment = Alignment(horizontal="center", wrap_text=True)
def row(r, vals, fmt="0.0000"):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.border = BOX
        if isinstance(v, float): c.number_format = fmt
r = 1
ws.cell(r, 1, "19대 선관위 분류기 통계 (구·시·군) — data19 대조와 후보별 비교").font = Font(bold=True, size=14); r += 1
ws.cell(r, 1, "원자료: 선관위 '분류기를 통과한 투표지수 / 분류된 투표지수 / 미분류 처리된 투표지수' (6명 후보별). 원자료는 아래 5번에 그대로 넣었다."); r += 2

ws.cell(r, 1, "1. data19 와 대조").font = B; r += 1
hdr(r, ["data19 열", "선관위 열", "일치 곳 수", "차이 합"]); r += 1
P = [("H  M1 (문재인 분류)", "C문", "M1"), ("I  H1 (홍준표 분류)", "C홍", "H1"), ("J  D1 (안철수 분류)", "C안", "D1"), ("K  M2 (문재인 미분류)", "U문", "M2"),
     ("L  H2 (홍준표 미분류)", "U홍", "H2"), ("M  D2 (안철수 미분류)", "U안", "D2"), ("N  U2", "U무효 (미분류 중 무효)", "U2"), ("F  vote_all", "T계 (분류기 통과 계)", "vote_all"), ("G  U_all", "U계 (미분류 계)", "U_all")]
for a, b, k in P:
    d = m[k].astype(float) - m[b.split()[0]]; row(r, [a, b, int((d == 0).sum()), int(d.sum())]); r += 1
ws.cell(r, 1, "249곳 모두 일치 (청주시흥덕구 = 흥덕 + 서원). data19 는 선관위 분류기 통계와 같은 자료다."); r += 2

ws.cell(r, 1, "2. 공개 최종득표보다 적은 이유").font = B; r += 1
ws.cell(r, 1, f"선관위 분류기 통과 투표지만 담았기 때문이다. 전국 문재인 공개 {int(x.문재인.sum()):,} vs 분류기 통과 {int(m.T문.sum()):,}, 홍준표 {int(x.홍준표.sum()):,} vs {int(m.T홍.sum()):,}."); r += 1
ws.cell(r, 1, "공개값보다 2% 이상 적은 14곳 (분류기를 거치지 않고 센 투표지가 많은 곳):"); r += 1
hdr(r, ["구·시·군", "공개 문재인", "분류기 문재인", "공개 홍준표", "분류기 홍준표", "부족 비율"]); r += 1
x2 = x.merge(m[["index", "T문", "T홍"]], on="index")
for _, v in x2[x2.rel >= .02].sort_values("rel", ascending=False).iterrows():
    row(r, [v.district, int(v.문재인), int(v.T문), int(v.홍준표), int(v.T홍), float(v.rel)]); ws.cell(r, 6).number_format = "0.0%"; r += 1
r += 1

ws.cell(r, 1, "3. 후보별 재확인(미분류)표 치우침 — data19 에 없던 유승민·심상정·기타 포함 (전국)").font = B; r += 1
hdr(r, ["후보", "분류표 득표율", "미분류 득표율", "비율의 비", "K (문재인 대비, 전국 합산)", "K (문재인 대비, 구·시·군 평균)", "K > 1 인 곳"]); r += 1
cand = [("문", "문재인"), ("홍", "홍준표"), ("안", "안철수"), ("유", "유승민"), ("심", "심상정"), ("기타", "기타 후보")]
Cv = g[[f"C{c}" for c, _ in cand]].sum().values; Uv = g[[f"U{c}" for c, _ in cand]].sum().values
for i, (c, nm) in enumerate(cand):
    o = (m[f"U{c}"] / m["U문"]) / (m[f"C{c}"] / m["C문"])
    row(r, [nm, Cv[i] / Cv.sum(), Uv[i] / Uv.sum(), (Uv[i] / Uv.sum()) / (Cv[i] / Cv.sum()), (Uv[i] / Uv[0]) / (Cv[i] / Cv[0]),
            float(o.replace(np.inf, np.nan).mean()), int((o > 1).sum())]); r += 1
ws.cell(r, 1, "보수 후보라도 유승민은 K < 1 (0.92) 이다. 홍준표 1.61, 기타 후보 2.56 으로 커서, 치우침은 이념보다 후보 지지층(고령층 등)의 표기 방식과 관련될 수 있다."); r += 1
ws.cell(r, 1, f"무효표 {int(g.T무효.sum()):,} 장은 모두 미분류로 분류된다 (미분류 {int(g.U계.sum()):,} 장의 {g.U무효.sum() / g.U계.sum():.1%})."); r += 2

ws.cell(r, 1, "4. 열 설명").font = B; r += 1
ws.cell(r, 1, "T = 분류기를 통과한 투표지, C = 분류된 투표지, U = 미분류 처리된 투표지. 문·홍·안·유·심·기타 = 후보별, 무효 = 무효표."); r += 2

ws.cell(r, 1, "5. 선관위 원자료 (시도 합계 + 구·시·군 250행)").font = B; r += 1
wsrc = openpyxl.load_workbook("nec19.xlsx").active
cols = ["시도", "구시군", "T계", "T유효", "T문", "T홍", "T안", "T유", "T심", "T기타", "T무효", "C계", "C문", "C홍", "C안", "C유", "C심", "C기타",
        "U계", "U유효", "U문", "U홍", "U안", "U유", "U심", "U기타", "U무효", "미분류율(%)"]
hdr(r, cols); r += 1
for src in wsrc.iter_rows(min_row=4, values_only=True):
    vals = [src[0], src[1]] + list(src[5:31]); row(r, vals, "0.00")
    if src[1] is None:
        for j in range(1, len(vals) + 1): ws.cell(r, j).font = B
    for j in range(3, 28): ws.cell(r, j).number_format = "#,##0"
    r += 1
ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 22
for j in range(3, 29): ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 13
ws.freeze_panes = None
lg = wb["점검_수정내역"]; n = lg.max_row + 1
lg.cell(n, 1, "nec19: 19대 선관위 분류기 통계(구·시·군 250행, 6명 후보) 원자료 추가. data19 의 분류·미분류 득표·무효·합계 9개 열이 249곳 모두 일치. 후보별 미분류/분류 비교(유승민 K 0.92, 홍준표 1.61, 기타 2.56).")
wb.save(F); print("saved rows", r)
