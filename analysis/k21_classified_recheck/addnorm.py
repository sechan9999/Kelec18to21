"""k18to21_수정본.xlsx 에 K정규성 탭 추가 (18·19대 구·시·군 K 정규성 검정, 이질성, 그림)"""
import shutil, openpyxl, pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image
F = "k18to21_수정본.xlsx"; shutil.copy(F, "k18to21_수정본_before_norm.xlsx")
S = pd.read_pickle("knorm.pkl"); T, Hh = S["T"], S["H"]
wb = openpyxl.load_workbook(F)
if "K정규성" in wb.sheetnames: del wb["K정규성"]
ws = wb.create_sheet("K정규성", wb.sheetnames.index("K정의_비교") + 1)
B = Font(bold=True); HF = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def hdr(r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = B; c.fill = HF; c.border = BOX; c.alignment = Alignment(horizontal="center", wrap_text=True)
def row(r, vals, fmt="0.0000"):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.border = BOX
        if isinstance(v, float): c.number_format = fmt
r = 1
ws.cell(r, 1, "18·19대 구·시·군 K 의 정규성").font = Font(bold=True, size=14); r += 1
ws.cell(r, 1, "K = (미분류 A/B) ÷ (분류 A/B). 18대 뉴스타파 251곳(박/문), 19대 선관위 개표단위 합계 250곳(각 후보/문재인). 자료: K18_K19_results.xlsx"); r += 2
ws.cell(r, 1, "1. 정규성 검정 (p < 0.05 이면 정규성 기각. Anderson-Darling 은 통계량이 5% 기준값보다 크면 기각)").font = B; r += 1
cols = ["n", "평균", "중앙값", "SD", "왜도", "첨도", "SW_p", "DAgostino_p", "JB_p", "AD_stat", "AD_5pct", "KS_p"]
hdr(r, ["K", "n", "평균", "중앙값", "SD", "왜도", "첨도", "Shapiro-Wilk p", "D'Agostino p", "Jarque-Bera p", "Anderson-Darling", "AD 5% 기준", "KS p", "판정 (5%)"]); r += 1
for k, t in T.iterrows():
    ok = t.SW_p >= .05 and t.AD_stat < t.AD_5pct
    row(r, [k, int(t.n)] + [float(t[c]) for c in cols[1:]] + ["기각 안 됨" if ok else "기각"])
    for j in (8, 9, 10, 13): ws.cell(r, j).number_format = "0.0000E+00" if t[cols[j - 2]] < 1e-4 else "0.0000"
    if ok: ws.cell(r, 14).font = Font(bold=True, color="1F7A1F")
    r += 1
r += 1
ws.cell(r, 1, "2. 구·시·군 K 가 모두 같은 값인가 (이질성)").font = B; r += 1
hdr(r, ["K", "n", "표본오차 중앙값 (log K)", "표준화 값 SD (기대 1)", "Q", "자유도", "I²", "구·시·군 간 SD (log K)", "표준화 값 SW p"]); r += 1
for k, h in Hh.iterrows():
    row(r, [k, int(h.n), float(h.표본오차_중앙값), float(h.z_SD), float(h.Q), int(h.df), float(h.I2), float(h.tau), float(h.z_SW_p)]); ws.cell(r, 7).number_format = "0.0%"; r += 1
ws.cell(r, 1, "표준화 값 = (구·시·군 log K − 전국 가중평균) ÷ 표본오차. 모든 곳의 참 K 가 같다면 SD 1 인 정규분포를 따른다."); r += 2
ws.cell(r, 1, "3. 결론").font = B; r += 1
for t in ["18대 K 는 오른쪽으로 약간 치우쳐 5% 수준에서 정규성이 기각되지만, log K 는 모든 검정에서 기각되지 않는다. 18대 K 는 로그정규분포에 맞는다.",
          "19대 홍/문 K 는 종 모양(평균 1.600, 중앙값 1.620)이지만 정규·로그정규 모두 5% 수준에서 기각된다. 가운데가 뾰족하고 1.0–1.2 쪽 꼬리가 길다.",
          "19대 안철수·유승민·심상정 K 도 정규성이 기각된다.",
          "표준화 값의 SD 가 18대 3.66, 19대 2.80 으로 1 보다 훨씬 크다. K 흩어짐의 87–93% 는 표본오차가 아니라 구·시·군 사이의 실제 차이다(I²).",
          "log K 와 분류표 보수 비율(R1)의 상관은 18대 0.46, 19대 0.44 다. K 분포는 서로 다른 값이 섞인 혼합분포라, 정규분포 여부만으로 이상 여부를 판단하기 어렵다."]:
    ws.cell(r, 1, "• " + t); r += 1
r += 1
img = Image("K18_K19_정규성.png"); img.width, img.height = 1150, 685; ws.add_image(img, f"A{r}")
ws.column_dimensions["A"].width = 18
for j in range(2, 15): ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 13
lg = wb["점검_수정내역"]; k = lg.max_row + 1
lg.cell(k, 1, "K정규성: 18·19대 구·시·군 K 와 log K 의 정규성 검정(Shapiro-Wilk·D'Agostino·Jarque-Bera·Anderson-Darling·KS), 이질성(Q·I²), 분포·QQ 그림. 18대는 로그정규, 19대는 둘 다 5% 기각, 두 선거 모두 I² 87–93%.")
wb.save(F); print("saved", r)
