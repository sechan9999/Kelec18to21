"""k18to21_수정본.xlsx 에 nec18 탭 추가: 뉴스타파 18대 분류기 운영결과(구·시·군 251행) 원자료 + data18 대조 + K 비교"""
import shutil, openpyxl, pandas as pd, numpy as np, statsmodels.formula.api as smf
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
F = "k18to21_수정본.xlsx"; shutil.copy(F, "k18to21_수정본_before_nec18.xlsx")
N = pd.read_pickle("nec18.pkl"); m = pd.read_pickle("nec18_merge.pkl")
wb = openpyxl.load_workbook(F)
if "nec18" in wb.sheetnames: del wb["nec18"]
ws = wb.create_sheet("nec18", wb.sheetnames.index("data18") + 1)
B = Font(bold=True); H = PatternFill("solid", fgColor="DDEBF7")
thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def hdr(r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = B; c.fill = H; c.border = BOX; c.alignment = Alignment(horizontal="center", wrap_text=True)
def row(r, vals, fmt="0.0000"):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.border = BOX
        if isinstance(v, float): c.number_format = fmt
r = 1
ws.cell(r, 1, "18대 분류기 운영결과 (뉴스타파 공개, 선관위 자료) — data18 대조").font = Font(bold=True, size=14); r += 1
ws.cell(r, 1, "구·시·군 251행 (부천 원미·소사·오정 별도). 총투표 = 수개표 + 분류기 개표, 분류기 개표 = 분류 + 미분류 가 251행 모두 성립. 원자료는 아래 5번."); r += 2

ws.cell(r, 1, "1. data18 열과 대조 (249곳, 부천 3개 구 합산)").font = B; r += 1
hdr(r, ["data18 열", "뉴스타파 열", "일치 곳 수", "차이 합 (data18 − 뉴스타파)"]); r += 1
for a, b in [("P1", "분류_박근혜"), ("M1", "분류_문재인"), ("P2", "미분류_박근혜"), ("M2", "미분류_문재인"), ("U2", "미분류_무효"), ("vote_all", "분류기개표_계"), ("U_all", "미분류_계")]:
    d = m[a].astype(float) - m[b]; row(r, [a, b, int((d == 0).sum()), int(d.sum())]); r += 1
r += 1

ws.cell(r, 1, "2. 차이 유형").font = B; r += 1
hdr(r, ["유형", "곳 수", "내용"]); r += 1
for c, t in [("A", "완전 일치. data18 = 분류기 개표분"), ("B", "후보별 합은 같고 분류↔미분류로 소수 표(최대 54표)가 옮겨짐 — 입력 차이"),
             ("C", "data18 이 수개표(부재자 등) 득표 일부·전부를 분류·미분류에 더함 (포함 비율 중앙값 92%)"), ("D", "기타 소수 차이 (최대 56표)")]:
    d = m[m.cat.str.startswith(c)]; row(r, [d.cat.iloc[0], len(d), t]); r += 1
ws.cell(r, 1, "공개 최종득표 = 뉴스타파 '계' (249곳 모두 확인). 공개값보다 적은 몫은 모두 수개표이며, 투표구가 빠진 곳은 없다."); r += 2

ws.cell(r, 1, "3. 수개표를 포함한 51곳 (data18 − 뉴스타파)").font = B; r += 1
hdr(r, ["시도", "구·시·군", "P1 차이", "M1 차이", "P2 차이", "M2 차이", "수개표 박근혜", "수개표 문재인"]); r += 1
for _, x in m[m.cat.str.startswith("C")].iterrows():
    row(r, [x.시도, x.district, int(x.dP1), int(x.dM1), int(x.dP2), int(x.dM2), int(x.수개표_박근혜), int(x.수개표_문재인)]); r += 1
r += 1

ws.cell(r, 1, "4. K 비교 (K = (미분류 박/문) ÷ (분류 박/문))").font = B; r += 1
hdr(r, ["자료", "n", "K 구·시·군 평균", "K 중앙값", "K 전국 합산", "적합식 절편", "적합식 기울기", "R²"]); r += 1
def fit(d, p1, m1, p2, m2):
    k = (d[p2] / d[m2]) / (d[p1] / d[m1]); t = pd.DataFrame({"R1": d[p1] / (d[p1] + d[m1]), "R2": d[p2] / (d[p2] + d[m2])}).astype(float)
    f = smf.ols("R2 ~ R1", t).fit()
    return [len(d), float(k.mean()), float(k.median()), float((d[p2].sum() / d[m2].sum()) / (d[p1].sum() / d[m1].sum())), float(f.params.Intercept), float(f.params.R1), float(f.rsquared)]
mm = m.copy()
for c in ["P1", "M1", "P2", "M2"]: mm[c] = mm[c].astype(float)
row(r, ["data18 (249)", *fit(mm, "P1", "M1", "P2", "M2")]); r += 1
row(r, ["뉴스타파, 부천 합산 (249)", *fit(mm, "분류_박근혜", "분류_문재인", "미분류_박근혜", "미분류_문재인")]); r += 1
row(r, ["뉴스타파 원자료 (251)", *fit(N, "분류_박근혜", "분류_문재인", "미분류_박근혜", "미분류_문재인")]); r += 1
for rr in range(r - 3, r): ws.cell(rr, 2).number_format = "0"
ws.cell(r, 1, "수개표 포함 여부와 무관하게 K 와 적합식은 거의 같다. 뉴스타파 251행 K 평균 1.481 은 기사의 18대 K ≈ 1.48 과 같다."); r += 2

ws.cell(r, 1, "5. 뉴스타파 원자료 (251행)").font = B; r += 1
cols = list(N.columns[:35])
hdr(r, cols); r += 1
for v in N[cols].itertuples(index=False):
    row(r, [None if (isinstance(a, float) and np.isnan(a)) else a for a in v], "#,##0")
    for j in range(31, 36): ws.cell(r, j).number_format = "0.00"
    r += 1
ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 16
for j in range(4, 36): ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 12
lg = wb["점검_수정내역"]; n = next((c.row for c in lg["A"] if isinstance(c.value, str) and c.value.startswith("nec18:")), lg.max_row + 1)
lg.cell(n, 1, "nec18: 뉴스타파 18대 분류기 운영결과(251행) 추가. data18 과 158곳 완전 일치, 40곳 소수 표 차이, 51곳은 data18 이 수개표(부재자 등)를 포함. 공개값과의 차이는 모두 수개표이며 투표구 누락은 없음. K 평균 data18 1.479 / 뉴스타파 1.481.")
wb.save(F); print("saved rows", r)
