"""K 정의·집계 방식별 비교 (18–21대), 18대 공개값 부족분 원인"""
import openpyxl, pandas as pd, numpy as np
num = lambda s: pd.to_numeric(s.astype(str).str.replace(",", "").str.strip(), errors="coerce")
wv = openpyxl.load_workbook("k18to21_수정본.xlsx", data_only=True)
def sheet(n):
    ws = wv[n]; h = [c.value.strip() if isinstance(c.value, str) else c.value for c in ws[1]]
    return pd.DataFrame([[c.value for c in ws[r]] for r in range(2, 251)], columns=h)
S = {"18대": (sheet("data18"), "P1", "M1", "P2", "M2", "K"), "19대": (sheet("data19"), "H1", "M1", "H2", "M2", "K"),
     "20대": (sheet("data20"), "Y1", "L1", "Y2", "L2", "K"), "21대": (sheet("data21"), "K1", "L1", "K2", "L2", "K")}
rows = []
for e, (d, c1, d1, c2, d2, kcol) in S.items():
    d = d.dropna(subset=[c1]).copy()
    if e == "20대": d = d[d.district != "제천시"]          # 제천: 분류표 합 = 총투표수, 공식 결과와 불일치 → 제외 (논문·대시보드 기준 248곳)
    for c in (c1, d1, c2, d2): d[c] = d[c].astype(float)
    orr = (d[c2] / d[d2]) / (d[c1] / d[d1])                         # 시트의 K 정의 = R2/R1, R = 보수/민주 (비의 비 = OR)
    ks = (d[c2] / (d[c2] + d[d2])) / (d[c1] / (d[c1] + d[d1]))      # 비율의 비 K = R_2/R_1
    rows.append(dict(선거=e, n=len(d), 시트K열_평균=float(pd.to_numeric(d[kcol], errors="coerce").mean()),
                     OR_구시군평균=orr.mean(), OR_중앙값=orr.median(), OR_전국합산=(d[c2].sum() / d[d2].sum()) / (d[c1].sum() / d[d1].sum()),
                     Kshare_구시군평균=ks.mean(), Kshare_전국합산=(d[c2].sum() / (d[c2] + d[d2]).sum()) / (d[c1].sum() / (d[c1] + d[d1]).sum())))
T = pd.DataFrame(rows).set_index("선거"); pd.set_option("display.width", 220); print(T.round(4).to_string())

# 18대 공개값 부족분: 구분별 공개 득표(국내부재자·재외·거소 등)와 비교
D = pd.read_excel("Kelec18to21/corrected_data/K18to21charts_corrected.xlsx", "18Data", header=None).iloc[3:]
D = D[D[1] != "합계"]
tot = D[(D[2] == "합계")]
spec = D[D[2].isin(["국내부재자투표", "재외투표", "국외부재자투표", "재외국민투표", "선상투표", "거소투표"]) & D[3].isna()]
print("18Data 특수 투표 행 이름:", sorted(D[2].dropna().astype(str).unique().tolist())[:0] or sorted(set(D[2].dropna().astype(str)) & {"국내부재자투표", "재외투표", "국외부재자투표", "재외국민투표", "선상투표", "거소투표"}))
sp = spec.groupby([0, 1]).agg(박=(6, lambda s: num(s).sum()), 문=(7, lambda s: num(s).sum())).reset_index()
x = pd.read_pickle("pe18_out.pkl")["x"]
tt = pd.DataFrame({"시도": tot[0].replace({"강원도": "강원특별자치도", "전라북도": "전북특별자치도"}).values, "구": tot[1].values})
sp["시도"] = sp[0].replace({"강원도": "강원특별자치도", "전라북도": "전북특별자치도", "세종특별자치시": "충청남도"}); sp["구"] = sp[1]
sp.loc[sp.구 == "세종특별자치시", "구"] = "세종시"; sp.loc[sp.구.astype(str).str.startswith("부천시"), "구"] = "부천시"
sp = sp.groupby(["시도", "구"], as_index=False)[["박", "문"]].sum()
m = x.merge(sp, on=["시도", "구"], how="left")
m["부족_박"] = -m.d박; m["부족_문"] = -m.d문
print("\n18대 부족분(공개 − pe18) vs 특수투표(부재자·재외 등) 득표:")
print("부족 박 합", int(m.부족_박.sum()), "| 특수투표 박 합", int(m.박.sum()), "| 부족 문 합", int(m.부족_문.sum()), "| 특수투표 문 합", int(m.문.sum()))
m["일치"] = (abs(m.부족_박 - m.박) <= 1) & (abs(m.부족_문 - m.문) <= 1)
print("구시군별로 부족분 = 특수투표 득표 인 곳:", int(m.일치.sum()), "/", len(m))
print(m[~m.일치][["district", "부족_박", "박", "부족_문", "문"]].head(15).to_string(index=False))
T.to_pickle("kcheck.pkl"); m.to_pickle("pe18_shortfall.pkl")
